"""Отрисовка полных отчётов в PDF (reportlab) и Excel (openpyxl)."""
import io
import os
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from data import MSP, dt, money, num, pct, pf, pretty

FONT_SOURCES = {
    'DejaVuSans.ttf': ['https://cdn.jsdelivr.net/npm/dejavu-fonts-ttf@2.37.3/ttf/DejaVuSans.ttf',
                       'https://raw.githubusercontent.com/dejavu-fonts/dejavu-fonts/version_2_37/ttf/DejaVuSans.ttf'],
    'DejaVuSans-Bold.ttf': ['https://cdn.jsdelivr.net/npm/dejavu-fonts-ttf@2.37.3/ttf/DejaVuSans-Bold.ttf',
                            'https://raw.githubusercontent.com/dejavu-fonts/dejavu-fonts/version_2_37/ttf/DejaVuSans-Bold.ttf'],
}
_FONTS_READY = {'ok': False}


def _font_bytes(name):
    """Шрифт с кириллицей: локальный кэш → хранилище проекта → публичный CDN (с сохранением в хранилище)."""
    local = f'/tmp/{name}'
    if os.path.exists(local) and os.path.getsize(local) > 100_000:
        return local
    import boto3
    import urllib.request
    s3 = boto3.client('s3', endpoint_url='https://bucket.poehali.dev', aws_access_key_id=os.environ['AWS_ACCESS_KEY_ID'],
                      aws_secret_access_key=os.environ['AWS_SECRET_ACCESS_KEY'])
    key = f'assets/fonts/{name}'
    body = None
    try:
        body = s3.get_object(Bucket='files', Key=key)['Body'].read()
    except Exception:
        for url in FONT_SOURCES[name]:
            try:
                with urllib.request.urlopen(url, timeout=6) as r:
                    body = r.read()
                if len(body) > 100_000:
                    s3.put_object(Bucket='files', Key=key, Body=body, ContentType='font/ttf')
                    break
            except Exception:
                body = None
    if not body:
        raise RuntimeError(f'Не удалось получить шрифт {name}')
    with open(local, 'wb') as f:
        f.write(body)
    return local


def ensure_fonts():
    if _FONTS_READY['ok']:
        return
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    pdfmetrics.registerFont(TTFont('DV', _font_bytes('DejaVuSans.ttf')))
    pdfmetrics.registerFont(TTFont('DVB', _font_bytes('DejaVuSans-Bold.ttf')))
    registerFontFamily('DV', normal='DV', bold='DVB', italic='DV', boldItalic='DVB')
    _FONTS_READY['ok'] = True

VIOLET = colors.HexColor('#6d28d9')
CYAN = colors.HexColor('#0891b2')
DARK = colors.HexColor('#111827')
GREY = colors.HexColor('#6b7280')
LIGHT = colors.HexColor('#f3f4f6')
LINE = colors.HexColor('#e5e7eb')
GREEN = colors.HexColor('#059669')
RED = colors.HexColor('#dc2626')
AMBER = colors.HexColor('#d97706')

ST = {
    'h1': ParagraphStyle('h1', fontName='DVB', fontSize=20, leading=25, textColor=DARK, spaceAfter=4),
    'h2': ParagraphStyle('h2', fontName='DVB', fontSize=13, leading=17, textColor=VIOLET, spaceBefore=12, spaceAfter=6),
    'sub': ParagraphStyle('sub', fontName='DV', fontSize=10, leading=14, textColor=GREY),
    'p': ParagraphStyle('p', fontName='DV', fontSize=9.5, leading=13.5, textColor=DARK, alignment=TA_LEFT),
    'small': ParagraphStyle('small', fontName='DV', fontSize=7.5, leading=10, textColor=GREY),
    'cell': ParagraphStyle('cell', fontName='DV', fontSize=8, leading=10, textColor=DARK),
    'cellb': ParagraphStyle('cellb', fontName='DVB', fontSize=8, leading=10, textColor=DARK),
    'kv': ParagraphStyle('kv', fontName='DVB', fontSize=14, leading=17, textColor=VIOLET),
    'kl': ParagraphStyle('kl', fontName='DV', fontSize=7.5, leading=9.5, textColor=GREY),
}


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('DV', 7)
    canvas.setFillColor(GREY)
    canvas.drawString(15 * mm, 10 * mm, 'МАТ-Лабс · Бизнес-аналитика · mat-labs.ru/analytics')
    canvas.drawRightString(195 * mm, 10 * mm, f'стр. {doc.page}')
    canvas.setStrokeColor(VIOLET)
    canvas.setLineWidth(2)
    canvas.line(15 * mm, 287 * mm, 195 * mm, 287 * mm)
    canvas.restoreState()


def doc_build(story, meta):
    ensure_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=16 * mm,
                            bottomMargin=16 * mm, title=meta.get('title', ''), author='МАТ-Лабс')
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def kpis(items, cols=4):
    cells = [[Paragraph(v, ST['kv']), Paragraph(l, ST['kl'])] for v, l in items]
    data, row = [], []
    for c in cells:
        row.append(c)
        if len(row) == cols:
            data.append(row)
            row = []
    if row:
        data.append(row + [''] * (cols - len(row)))
    t = Table(data, colWidths=[180 * mm / cols] * cols)
    t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), LIGHT), ('BOX', (0, 0), (-1, -1), 0, colors.white),
                           ('INNERGRID', (0, 0), (-1, -1), 3, colors.white), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                           ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
                           ('LEFTPADDING', (0, 0), (-1, -1), 8)]))
    return t


def table(head, body, widths, align_right=()):
    data = [[Paragraph(h, ST['cellb']) for h in head]] + [[Paragraph(str(x), ST['cell']) for x in r] for r in body]
    t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    style = [('BACKGROUND', (0, 0), (-1, 0), LIGHT), ('LINEBELOW', (0, 0), (-1, -1), 0.3, LINE),
             ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 3),
             ('BOTTOMPADDING', (0, 0), (-1, -1), 3), ('LEFTPADDING', (0, 0), (-1, -1), 4),
             ('RIGHTPADDING', (0, 0), (-1, -1), 4)]
    for c in align_right:
        style.append(('ALIGN', (c, 0), (c, -1), 'RIGHT'))
    t.setStyle(TableStyle(style))
    return t


def bar_chart(labels, values, title, w=180, h=60, color=VIOLET):
    d = Drawing(w * mm, h * mm)
    ch = VerticalBarChart()
    ch.x, ch.y = 10 * mm, 12 * mm
    ch.width, ch.height = (w - 15) * mm, (h - 20) * mm
    ch.data = [[float(v or 0) for v in values]]
    ch.categoryAxis.categoryNames = [str(x) for x in labels]
    ch.categoryAxis.labels.fontName = 'DV'
    ch.categoryAxis.labels.fontSize = 6.5
    ch.categoryAxis.labels.angle = 0 if len(labels) <= 8 else 30
    ch.valueAxis.visible = False
    ch.valueAxis.valueMin = 0
    ch.bars[0].fillColor = color
    ch.bars[0].strokeColor = None
    ch.barSpacing = 2
    d.add(ch)
    d.add(String(0, (h - 5) * mm, title, fontName='DVB', fontSize=8.5, fillColor=DARK))
    return d


def pie_chart(items, title, w=85, h=55):
    items = [(l, v) for l, v in items if v]
    d = Drawing(w * mm, h * mm)
    p = Pie()
    p.x, p.y = 4 * mm, 4 * mm
    p.width = p.height = (h - 14) * mm
    p.data = [float(v) for _, v in items] or [1]
    p.labels = None
    pal = [VIOLET, CYAN, colors.HexColor('#a78bfa'), colors.HexColor('#67e8f9'), GREY]
    for i in range(len(p.data)):
        p.slices[i].fillColor = pal[i % len(pal)]
        p.slices[i].strokeColor = colors.white
    d.add(p)
    tot = sum(p.data) or 1
    for i, (l, v) in enumerate(items):
        y = (h - 16 - i * 6) * mm
        d.add(String((h - 6) * mm, y, f'■ {l}: {v / tot * 100:.0f}%'.replace('.', ','), fontName='DV', fontSize=7,
                     fillColor=pal[i % len(pal)]))
    d.add(String(0, (h - 5) * mm, title, fontName='DVB', fontSize=8.5, fillColor=DARK))
    return d


def cover(story, kind, title, subtitle, order):
    story.append(Paragraph(kind.upper(), ParagraphStyle('k', fontName='DVB', fontSize=9, textColor=CYAN, spaceAfter=4)))
    story.append(Paragraph(title, ST['h1']))
    story.append(Paragraph(subtitle, ST['sub']))
    story.append(Spacer(1, 3 * mm))
    meta = f'Заказ №{order["id"]} · подготовлено {date.today().strftime("%d.%m.%Y")} для: {order.get("company") or order.get("name")}'
    story.append(Paragraph(meta, ST['small']))
    story.append(Spacer(1, 5 * mm))


SOURCE = ('Источники: открытые данные ФНС России за 2025 год (доходы и расходы по бухгалтерской отчётности, среднесписочная '
          'численность, специальные налоговые режимы, недоимки), единый реестр субъектов МСП на 10.08.2026, ЕГРЮЛ на дату '
          'подготовки отчёта, реестр получателей поддержки МСП. Компании без бухгалтерской отчётности в ФНС учитываются в '
          'числе игроков, но не в объёме выручки. Отчёт носит информационный характер.')


# ── Рынок ───────────────────────────────────────────────────────────────────

def market_pdf(d, conclusions, order, kind='Отчёт о рынке'):
    ensure_fonts()
    m, nat = d['metrics'], d['national']
    s = []
    cover(s, kind, d['okved_name'], f'{d["region_name"]} · ОКВЭД {d["okved"] if d["okved"] != "ALL" else "все"} · '
          f'данные за 2025 год · {num(m["companies"])} компаний', order)
    s.append(kpis([
        (money(m['revenue']), 'объём рынка, выручка 2025'),
        (num(m['companies']), f'компаний, из них с отчётностью {num(m["with_revenue"])}'),
        (money(m['median_revenue']), 'медианная выручка'),
        (num(m['staff']), 'сотрудников'),
        (pf(m["profitable"], m["with_revenue"]), 'с доходами выше расходов'),
        (pf(m["top10_share"], d=0), f'доля топ-10, HHI {float(m["hhi"]):.0f}'),
        (pf(m["new_3y"], m["companies"]), 'новые игроки 2024–2026'),
        (pf(m["with_debt"], m["companies"]), f'с недоимкой, {money(m["debt_sum"])}'),
    ]))
    if nat and d['region'] != 'RU':
        s.append(Spacer(1, 2 * mm))
        s.append(Paragraph(f'Доля региона в рынке России: <b>{pf(m["revenue"], nat["revenue"])}</b> выручки и '
                           f'<b>{pf(m["companies"], nat["companies"])}</b> компаний. Медианная выручка по России — '
                           f'{money(nat["median_revenue"])}.', ST['p']))
    s.append(Paragraph('Выводы аналитика', ST['h2']))
    for i, c in enumerate(conclusions, 1):
        s.append(Paragraph(f'<b>{i}.</b> {c}', ST['p']))
        s.append(Spacer(1, 1.5 * mm))

    s.append(Paragraph('Лидеры рынка — топ-50 по выручке', ST['h2']))
    tot = float(m['revenue'] or 0) or 1
    body = [[i + 1, f'{pretty(t["name"])}<br/><font size="6.5" color="#6b7280">ИНН {t["inn"]} · {t.get("city") or ""}</font>',
             money(t['revenue']), pf(t["revenue"], tot),
             money(t['revenue'] - t['expenses']) if t.get('expenses') is not None else '—',
             t['staff'] if t.get('staff') is not None else '—', MSP.get(t['msp_cat'], '')]
            for i, t in enumerate(d['top'])]
    s.append(table(['#', 'Компания', 'Выручка', 'Доля', 'Дох. − расх.', 'Штат', 'Размер'], body,
                   [8, 70, 24, 13, 24, 13, 28], align_right=(2, 3, 4, 5)))

    s.append(PageBreak())
    s.append(Paragraph('Структура рынка', ST['h2']))
    if d['segments']:
        body = [[x['code'], x['name'] or '—', num(x['companies']), money(x['revenue']),
                 pf(x["revenue"], tot), num(x.get('staff') or 0)] for x in d['segments']]
        s.append(table(['Код', 'Сегмент', 'Компаний', 'Выручка', 'Доля', 'Штат'], body, [14, 88, 18, 24, 14, 22],
                       align_right=(2, 3, 4, 5)))
        s.append(Spacer(1, 4 * mm))
    charts = Table([[pie_chart([('Микро', m['micro']), ('Малые', m['small']), ('Средние', m['medium']),
                                ('Крупные и прочие', m['non_msp'])], 'Компании по размеру'),
                     bar_chart([x['label'] for x in d['size_dist']], [x['revenue'] for x in d['size_dist']],
                               'Выручка по размеру компаний', w=92, h=55)]], colWidths=[88 * mm, 92 * mm])
    s.append(charts)
    if d['cohorts']:
        s.append(Spacer(1, 3 * mm))
        s.append(bar_chart([x['y'] for x in d['cohorts']], [x['n'] for x in d['cohorts']],
                           'Приток новых игроков: год входа в реестр МСП', color=CYAN, h=50))

    s.append(Paragraph('География', ST['h2']))
    body = [[i + 1, g['name'], num(g['companies']), money(g['revenue']), money(g['median']) if g.get('median') else '—'] for i, g in enumerate(d['geo'])]
    s.append(table(['#', 'Регион', 'Компаний', 'Выручка', 'Медиана'], body, [8, 70, 30, 36, 36], align_right=(2, 3, 4)))

    s.append(Paragraph('Финансы и риски рынка', ST['h2']))
    margin = pf(float(m['revenue']) - float(m['expenses']), float(m['revenue']))
    body = [['Совокупные расходы', money(m['expenses'])],
            ['Совокупный результат (доходы − расходы)', money(float(m['revenue']) - float(m['expenses']))],
            ['Рентабельность рынка', margin],
            ['Выручка на сотрудника', money(float(m['revenue']) / float(m['staff'])) if m['staff'] else '—'],
            ['Доля компаний на УСН и АУСН', pf(m["usn"], m["companies"])],
            ['Компаний с налоговой недоимкой', f'{num(m["with_debt"])} ({pf(m["with_debt"], m["companies"])})'],
            ['Сумма недоимок', money(m['debt_sum'])]]
    s.append(table(['Показатель', 'Значение'], body, [120, 60], align_right=(1,)))
    if d['grants']:
        s.append(Paragraph('Получатели господдержки 2024–2026', ST['h2']))
        body = [[pretty(g['name']), g['inn'], money(g['total_amount']), dt(g['last_date']), g['source']]
                for g in d['grants'][:20]]
        s.append(table(['Компания', 'ИНН', 'Сумма', 'Последняя', 'Источник'], body, [60, 24, 24, 22, 50], align_right=(2,)))
    s.append(Spacer(1, 6 * mm))
    s.append(Paragraph(f'Полный список — {num(len(d["all"]))} компаний с показателями и контактами (если найдены) — '
                       f'в приложенном файле Excel.', ST['p']))
    s.append(Spacer(1, 4 * mm))
    s.append(Paragraph(SOURCE, ST['small']))
    return doc_build(s, {'title': f'Отчёт о рынке: {d["okved_name"]} — {d["region_name"]}'})


def xl_sheet(wb, title, head, rows_, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='6D28D9')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows_:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def money_fmt(ws, cols):
    for col in cols:
        for c in ws[get_column_letter(col)][1:]:
            c.number_format = '#,##0'


XL_LIMIT = 3000
CSV_HEAD = ['Компания', 'ИНН', 'Код региона', 'Город', 'ОКВЭД', 'Вид деятельности', 'Выручка 2025, ₽', 'Расходы 2025, ₽',
            'Доходы − расходы, ₽', 'Штат', 'Размер', 'В реестре МСП с', 'Налоговый режим', 'Недоимка, ₽', 'Руководитель',
            'Сайт', 'Телефоны', 'Email']


def company_row(x, names):
    return [pretty(x['name']), x['inn'], x['region'], x.get('city') or '', x['okved'],
            names.get(x['okved'][:5]) or names.get(x['okved'][:4]) or '',
            x['revenue'], x['expenses'],
            (x['revenue'] - x['expenses']) if x['revenue'] is not None and x['expenses'] is not None else None,
            x['staff'], MSP.get(x['msp_cat'], ''), x['msp_year'] or None, x['tax_mode'], x['tax_debt'] or 0,
            x.get('head') or '', x.get('website') or '', x.get('phones') or '', x.get('emails') or '']


def full_csv(d):
    import csv
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=';')
    w.writerow(CSV_HEAD)
    names = d['okved_names']
    for x in d['all']:
        w.writerow(['' if v is None else v for v in company_row(x, names)])
    return ('\ufeff' + buf.getvalue()).encode('utf-8')


def market_xlsx(d, conclusions, order):
    m = d['metrics']
    wb = Workbook(write_only=False)
    summary = [['Отрасль', d['okved_name']], ['Код ОКВЭД', d['okved']], ['Регион', d['region_name']], ['Период', '2025 год'],
               ['Компаний', m['companies']], ['С отчётностью о выручке', m['with_revenue']],
               ['Объём рынка (выручка), ₽', int(m['revenue'])], ['Расходы, ₽', int(m['expenses'])],
               ['Медианная выручка, ₽', int(m['median_revenue'])], ['Сотрудников', int(m['staff'])],
               ['С доходами выше расходов, %', pct(m['profitable'], m['with_revenue'])],
               ['Доля топ-10, %', float(m['top10_share'])], ['Индекс HHI', float(m['hhi'])],
               ['Микро', m['micro']], ['Малые', m['small']], ['Средние', m['medium']], ['Крупные и прочие', m['non_msp']],
               ['Новые 2024–2026', m['new_3y']], ['С недоимкой', m['with_debt']], ['Сумма недоимок, ₽', int(m['debt_sum'])],
               ['', ''], ['Выводы', '']] + [[f'{i}.', c] for i, c in enumerate(conclusions, 1)] + \
              [['', ''], ['Заказ', f'№{order["id"]} от {date.today().strftime("%d.%m.%Y")}'], ['Источник', SOURCE]]
    ws = xl_sheet(wb, 'Сводка', ['Показатель', 'Значение'], summary, [34, 110], first=True)
    for r in ws.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical='top')
        if isinstance(r[1].value, int) and r[1].value > 9999:
            r[1].number_format = '#,##0'
    names = d['okved_names']
    big = len(d['all']) > XL_LIMIT
    rows_ = [company_row(x, names) for x in d['all'][:XL_LIMIT]]
    ws = xl_sheet(wb, f'Топ-{num(XL_LIMIT)}' if big else 'Все компании', CSV_HEAD, rows_,
                  [42, 13, 8, 16, 10, 36, 16, 16, 16, 8, 14, 10, 10, 12, 26, 22, 24, 30])
    if not big:
        money_fmt(ws, [7, 8, 9, 14])
    else:
        ws.insert_rows(1)
        ws['A1'] = f'Первые {num(XL_LIMIT)} компаний по выручке. Полный список ({num(len(d["all"]))}) — в файле CSV в этом же архиве.'
        ws['A1'].font = Font(bold=True, color='6D28D9')
        ws.freeze_panes = 'A3'
    rows_ = [[x['code'], x['name'], x['companies'], int(x['revenue']), int(x.get('staff') or 0)] for x in d['segments']]
    ws = xl_sheet(wb, 'Сегменты', ['Код', 'Сегмент', 'Компаний', 'Выручка, ₽', 'Штат'], rows_, [10, 70, 12, 18, 12])
    money_fmt(ws, [4])
    rows_ = [[g['name'], g['region'], g['companies'], int(g['revenue']), int(g.get('median') or 0)] for g in d['geo']]
    ws = xl_sheet(wb, 'Регионы', ['Регион', 'Код', 'Компаний', 'Выручка, ₽', 'Медиана, ₽'], rows_, [30, 8, 12, 18, 16])
    money_fmt(ws, [4, 5])
    if d['grants']:
        rows_ = [[pretty(g['name']), g['inn'], int(g['total_amount'] or 0), str(g['last_date'] or '')[:10], g['source']] for g in d['grants']]
        ws = xl_sheet(wb, 'Господдержка', ['Компания', 'ИНН', 'Сумма 2024–2026, ₽', 'Последняя выплата', 'Источник'], rows_, [42, 13, 18, 14, 34])
        money_fmt(ws, [3])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Компания / проверка ─────────────────────────────────────────────────────

STATUS = {'ACTIVE': 'Действующая', 'LIQUIDATING': 'Ликвидируется', 'LIQUIDATED': 'Ликвидирована', 'BANKRUPT': 'Банкрот',
          'REORGANIZING': 'Реорганизация'}
LEVEL = {'low': ('Низкий риск', GREEN), 'mid': ('Средний риск', AMBER), 'high': ('Высокий риск', RED)}
FLV = {'high': ('✖', '#dc2626'), 'mid': ('▲', '#d97706'), 'good': ('✔', '#059669')}


def reg_date(ms):
    return datetime.utcfromtimestamp(ms / 1000).strftime('%d.%m.%Y') if ms else '—'


def risk_block(c):
    r = c['risk']
    label, col = LEVEL[r['level']]
    head = Table([[Paragraph(f'<font size="26" color="{col.hexval()}"><b>{r["score"]}</b></font><font size="9" color="#6b7280"> / 100</font>', ST['p']),
                   Paragraph(f'<font color="{col.hexval()}"><b>{label}</b></font><br/><font size="7.5" color="#6b7280">'
                             f'индекс риска МАТ-Лабс по открытым данным: чем выше, тем рискованнее</font>', ST['p'])]],
                 colWidths=[35 * mm, 145 * mm])
    head.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), LIGHT), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                              ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
    out = [head, Spacer(1, 2 * mm)]
    for lvl, text in r['factors'] or [('good', 'Существенных факторов риска в открытых данных не найдено')]:
        sym, color = FLV[lvl]
        out.append(Paragraph(f'<font color="{color}">{sym}</font>  {text}', ST['p']))
    return out


def company_pdf(c, order, check_only=False):
    ensure_fonts()
    s = []
    title = c['short'] or pretty(c['name'])
    cover(s, 'Проверка контрагента' if check_only else 'Отчёт о компании', title,
          f'ИНН {c["inn"]}' + (f' · ОГРН {c["ogrn"]}' if c['ogrn'] else '') + f' · {c["region_name"]}', order)
    s.append(Paragraph('Проверка контрагента', ST['h2']))
    s += risk_block(c)
    s.append(Spacer(1, 3 * mm))
    body = [['Статус в ЕГРЮЛ', STATUS.get(c['status'], c['status'] or 'нет данных')],
            ['Дата регистрации', reg_date(c['registered'])],
            ['Налоговая недоимка', money(c['tax_debt']) if c['tax_debt'] else 'нет'],
            ['Доходы − расходы 2025', money(c['profit']) if c['profit'] is not None else '—'],
            ['Налоговый режим', c['tax_mode'] or '—'], ['Категория МСП', c['msp'] or '—']]
    s.append(table(['Проверка', 'Результат'], body, [70, 110]))

    s.append(Paragraph('Реквизиты', ST['h2']))
    body = [['Полное наименование', c['full_name']], ['ИНН / КПП', f'{c["inn"]} / {c["kpp"] or "—"}'],
            ['ОГРН', c['ogrn'] or '—'], ['Адрес', c['address'] or '—'],
            [c['head_post'] or 'Руководитель', c['head'] or '—'],
            ['Основной вид деятельности', f'{c["okved"]} {c["okved_name"]}'],
            ['Уставный капитал', money(c['capital']) if c['capital'] is not None else '—'],
            ['Филиалов', c['branches'] if c['branches'] is not None else '—']]
    if c['founders']:
        body.append(['Учредители', '<br/>'.join(f'{f["name"]}' + (f' — {f["share"]}%' if f.get('share') else '') for f in c['founders'])])
    s.append(table(['Поле', 'Значение'], body, [50, 130]))
    if check_only:
        s.append(Spacer(1, 6 * mm))
        s.append(Paragraph(SOURCE + ' Индекс риска не является юридическим заключением.', ST['small']))
        return doc_build(s, {'title': f'Проверка контрагента: {title}'})

    s.append(Paragraph('Финансы 2025', ST['h2']))
    p = c['peers'] or {}
    s.append(kpis([
        (money(c['revenue']) if c['revenue'] is not None else '—', 'выручка'),
        (money(c['profit']) if c['profit'] is not None else '—', 'доходы минус расходы'),
        (pf(c["margin"]) if c['margin'] is not None else '—',
         f'рентабельность; медиана отрасли {pf(p["med_margin"] * 100)}' if p.get('med_margin') is not None else 'рентабельность'),
        (str(c['staff']) if c['staff'] is not None else '—',
         f'сотрудников; {money(c["rev_per_staff"])} на чел.' if c['rev_per_staff'] else 'сотрудников'),
    ]))
    if p:
        s.append(Spacer(1, 3 * mm))
        s.append(Paragraph(f'Место по выручке: <b>{p["rank"]}</b> из {num(p["n"])} в отрасли «{c["okved2_name"]}» в регионе, '
                           f'<b>{num(p["rank_ru"])}</b> из {num(p["n_ru"])} по России. Выручка в '
                           + (f'{float(c["revenue"]) / float(p["med"]):.1f}'.replace('.', ',') +
                              f' раза выше медианы отрасли в регионе ({money(p["med"])}).' if p.get('med') and c.get('revenue') else ''), ST['p']))
    if c['grant']:
        g = c['grant']
        s.append(Paragraph('Господдержка 2024–2026', ST['h2']))
        s.append(Paragraph(f'Получено <b>{money(g["total_amount"])}</b> за {g["payments"]} выплат; последняя — '
                           f'{money(g["last_amount"])} от {dt(g["last_date"])}. {g["source_full"]}.', ST['p']))
    ct = c['contacts']
    if any(ct.values()):
        s.append(Paragraph('Контакты', ST['h2']))
        s.append(table(['Канал', 'Значение'], [['Сайт', ct['website'] or '—'], ['Телефоны', ct['phones'] or '—'],
                                               ['Email', ct['emails'] or '—']], [40, 140]))
    if c['competitors']:
        s.append(Paragraph('Ближайшие конкуренты — та же отрасль и регион', ST['h2']))
        body = [[f'{pretty(x["name"])}<br/><font size="6.5" color="#6b7280">ИНН {x["inn"]}</font>', money(x['revenue']),
                 money(x['revenue'] - x['expenses']) if x.get('expenses') is not None else '—',
                 x['staff'] if x.get('staff') is not None else '—', MSP.get(x['msp_cat'], '')] for x in c['competitors'][:20]]
        s.append(table(['Компания', 'Выручка', 'Дох. − расх.', 'Штат', 'Размер'], body, [82, 26, 26, 16, 30], align_right=(1, 2, 3)))
    s.append(Spacer(1, 6 * mm))
    s.append(Paragraph(SOURCE + ' Индекс риска не является юридическим заключением.', ST['small']))
    return doc_build(s, {'title': f'Отчёт о компании: {title}'})


def company_xlsx(c, order):
    wb = Workbook()
    rows_ = [['Наименование', c['full_name']], ['ИНН', c['inn']], ['ОГРН', c['ogrn']], ['КПП', c['kpp']],
             ['Статус', STATUS.get(c['status'], c['status'])], ['Регистрация', reg_date(c['registered'])],
             ['Адрес', c['address']], [c['head_post'] or 'Руководитель', c['head']],
             ['ОКВЭД', f'{c["okved"]} {c["okved_name"]}'], ['Выручка 2025, ₽', c['revenue']], ['Расходы 2025, ₽', c['expenses']],
             ['Доходы − расходы, ₽', c['profit']], ['Рентабельность, %', c['margin']], ['Штат', c['staff']],
             ['Налоговый режим', c['tax_mode']], ['Недоимка, ₽', c['tax_debt']], ['Индекс риска', c['risk']['score']],
             ['Сайт', c['contacts']['website']], ['Телефоны', c['contacts']['phones']], ['Email', c['contacts']['emails']]]
    rows_ += [['Фактор риска', t] for _, t in c['risk']['factors']]
    ws = xl_sheet(wb, 'Компания', ['Поле', 'Значение'], rows_, [28, 90], first=True)
    for r in ws.iter_rows(min_row=2):
        if isinstance(r[1].value, int) and abs(r[1].value) > 9999:
            r[1].number_format = '#,##0'
    rows_ = [[pretty(x['name']), x['inn'], x.get('city') or '', x['revenue'], x['expenses'],
              (x['revenue'] - x['expenses']) if x.get('expenses') is not None else None, x['staff'], MSP.get(x['msp_cat'], ''),
              x['tax_debt'] or 0, x.get('website') or '', x.get('phones') or '', x.get('emails') or ''] for x in c['competitors']]
    ws = xl_sheet(wb, 'Конкуренты', ['Компания', 'ИНН', 'Город', 'Выручка, ₽', 'Расходы, ₽', 'Доходы − расходы, ₽', 'Штат', 'Размер',
                                     'Недоимка, ₽', 'Сайт', 'Телефоны', 'Email'], rows_, [42, 13, 16, 16, 16, 16, 8, 14, 12, 22, 24, 30])
    money_fmt(ws, [4, 5, 6, 9])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
